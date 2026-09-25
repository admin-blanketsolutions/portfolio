"""Untrusted bytes in, typed trial balance out.

The result carries control totals (line count, sum of debit closing balances,
sum of credit closing balances). The database recomputes them from the rows it
actually received and refuses to close the import on any mismatch, so a
parser bug or a tampered payload cannot silently change the TB.
"""

from __future__ import annotations

import codecs
import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Iterable, Iterator

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from . import PARSER_VERSION
from .amounts import fmt, parse_cell
from .errors import ParseError
from .layout import TOTAL_ROW, Layout, detect
from .limits import Limits
from .normalize import ascii_digits, clean_display, looks_like_formula, normalize_for_matching
from .ooxml_guard import inspect_xlsx

FORMATS = ("csv", "xlsx")
_AMOUNT_ROLES = ("opening", "opening_debit", "opening_credit", "period_debit", "period_credit",
                 "closing", "closing_debit", "closing_credit")
_MAX_CODE = 64
_MAX_NAME = 512


@dataclass
class Line:
    line_no: int
    source_row: int
    code: str
    name: str
    name_norm: str
    opening: Decimal | None
    period_debit: Decimal | None
    period_credit: Decimal | None
    closing: Decimal
    had_formula: bool

    def to_dict(self) -> dict[str, object]:
        return {
            "line_no": self.line_no, "source_row": self.source_row, "code": self.code, "name": self.name,
            "name_norm": self.name_norm, "opening": fmt(self.opening), "period_debit": fmt(self.period_debit),
            "period_credit": fmt(self.period_credit), "closing": fmt(self.closing), "had_formula": self.had_formula,
        }


@dataclass
class Warnings:
    limit: int
    items: dict[str, dict[str, object]] = field(default_factory=dict)

    def add(self, code: str, message: str, row: int | None = None) -> None:
        w = self.items.setdefault(code, {"code": code, "message": message, "count": 0, "rows": []})
        w["count"] = int(w["count"]) + 1  # type: ignore[call-overload]
        rows = w["rows"]
        assert isinstance(rows, list)
        if row is not None and len(rows) < self.limit:
            rows.append(row)

    def to_list(self) -> list[dict[str, object]]:
        return list(self.items.values())


def parse_bytes(data: bytes, fmt_name: str, *, limits: Limits = Limits(), sheet: str | None = None,
                decimal_separator: str | None = None) -> dict[str, object]:
    if fmt_name not in FORMATS:
        raise ParseError("unsupported_format", "Only .xlsx and .csv trial balances are accepted.")
    if len(data) > limits.max_file_bytes:
        raise ParseError("file_too_large", "The file exceeds the maximum upload size.")
    if not data:
        raise ParseError("empty", "The file is empty.")
    if decimal_separator not in (None, ".", ","):
        raise ParseError("invalid_option", "decimal_separator must be '.' or ','.")
    warnings = Warnings(limits.max_warning_rows)
    if fmt_name == "xlsx":
        sheet_name, rows = _xlsx_rows(data, limits, sheet, warnings)
    else:
        sheet_name, rows = None, _csv_rows(data, limits, warnings)
    return _build(rows, limits, decimal_separator, warnings, fmt_name, sheet_name)


# ---------------------------------------------------------------------------
# Readers: both yield (row_no, values, formula_flags)
# ---------------------------------------------------------------------------
Row = tuple[int, tuple[object, ...], tuple[bool, ...]]


def _xlsx_rows(data: bytes, limits: Limits, sheet: str | None, warnings: Warnings) -> tuple[str, list[Row]]:
    inspect_xlsx(data, limits)
    try:
        values_wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        formula_wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ParseError:
        raise
    except Exception as exc:  # openpyxl raises a zoo of exception types on malformed input
        raise ParseError("malformed", "The workbook could not be read.") from exc

    try:
        visible = [ws for ws in values_wb.worksheets
                   if isinstance(ws, ReadOnlyWorksheet) and getattr(ws, "sheet_state", "visible") == "visible"]
        hidden = len(values_wb.worksheets) - len(visible)
        if hidden:
            warnings.add("hidden_sheets_ignored", "Hidden sheets were ignored.")
        if sheet is not None:
            chosen = [ws for ws in visible if ws.title == sheet]
            if not chosen:
                raise ParseError("sheet_not_found", "The requested sheet does not exist or is hidden.")
            candidates = chosen
        else:
            candidates = visible
        if not candidates:
            raise ParseError("no_sheet", "The workbook has no visible worksheet.")

        last_error: ParseError | None = None
        for ws in candidates:
            fws = formula_wb[ws.title]
            head = list(_xlsx_iter(ws, fws, limits, limit=limits.header_scan_rows))
            try:
                detect([r[1] for r in head])
            except ParseError as exc:
                last_error = exc
                continue
            return ws.title, list(_xlsx_iter(ws, fws, limits))
        assert last_error is not None
        raise last_error
    finally:
        values_wb.close()
        formula_wb.close()


def _xlsx_iter(ws: ReadOnlyWorksheet, fws: ReadOnlyWorksheet, limits: Limits, limit: int | None = None) -> Iterator[Row]:
    # Read-only mode otherwise trusts the file's <dimension> tag and silently
    # stops there: a workbook that under-declares its size would have trailing
    # TB lines dropped without any error. Parse the actual rows instead (gaps
    # are filled with empty rows, so the row limit and row numbers stay true).
    ws.reset_dimensions()
    fws.reset_dimensions()
    vals = ws.iter_rows(max_col=limits.max_cols, values_only=True)
    forms = fws.iter_rows(max_col=limits.max_cols, values_only=True)
    for i, (v, f) in enumerate(zip(vals, forms), start=1):
        if limit is not None and i > limit:
            return
        if i > limits.max_rows:
            raise ParseError("too_many_rows", f"The sheet has more than {limits.max_rows} rows.")
        flags = tuple(isinstance(x, str) and x.startswith("=") for x in f)
        yield i, tuple(v), flags


def _decode_csv(data: bytes, warnings: Warnings) -> str:
    if data.startswith(codecs.BOM_UTF8):
        return data[3:].decode("utf-8", errors="strict")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    if b"\x00" in data:
        raise ParseError("malformed", "The CSV file contains NUL bytes.")
    try:
        return data.decode("utf-8", errors="strict")
    except UnicodeDecodeError:
        warnings.add("legacy_encoding", "The CSV is not UTF-8; it was read as Windows-1256 (Arabic).")
        return data.decode("cp1256", errors="strict")


def _csv_rows(data: bytes, limits: Limits, warnings: Warnings) -> list[Row]:
    try:
        text = _decode_csv(data, warnings)
    except UnicodeDecodeError as exc:
        raise ParseError("encoding", "The CSV encoding could not be determined; save it as UTF-8.") from exc
    sample = text[:16384]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    csv.field_size_limit(limits.max_cell_chars * 4)
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    rows: list[Row] = []
    try:
        for i, rec in enumerate(reader, start=1):
            if i > limits.max_rows:
                raise ParseError("too_many_rows", f"The file has more than {limits.max_rows} rows.")
            if len(rec) > limits.max_cols:
                rec = rec[: limits.max_cols]
            values = tuple(x if x.strip() != "" else None for x in rec)
            rows.append((i, values, tuple(False for _ in values)))
    except csv.Error as exc:
        raise ParseError("malformed", "The CSV file is malformed.") from exc
    return rows


# ---------------------------------------------------------------------------
# Row interpretation
# ---------------------------------------------------------------------------
def _cell(values: tuple[object, ...], idx: int | None) -> object:
    if idx is None or idx >= len(values):
        return None
    return values[idx]


def _text(value: object, limits: Limits, row: int, col: str) -> tuple[str, bool]:
    if value is None:
        return "", False
    if isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    elif isinstance(value, float):
        raw = repr(value)
    else:
        raw = str(value)
    if len(raw) > limits.max_cell_chars:
        raise ParseError("cell_too_long", f"Row {row}, column {col}: the cell text is too long.", row)
    return clean_display(raw)


def _build(rows: Iterable[Row], limits: Limits, decimal_separator: str | None, warnings: Warnings,
           fmt_name: str, sheet_name: str | None) -> dict[str, object]:
    rows = list(rows)
    layout: Layout = detect([r[1] for r in rows[: limits.header_scan_rows]], first_row_no=rows[0][0] if rows else 1)
    cols = layout.columns
    letters = layout.letters()
    lines: list[Line] = []
    seen_codes: dict[str, int] = {}
    footer: dict[str, Decimal | None] | None = None

    for row_no, values, formulas in rows:
        if row_no <= layout.header_row:
            continue
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        for role in ("code", "name"):
            idx = cols[role]
            if idx < len(formulas) and formulas[idx] and _cell(values, idx) is None:
                raise ParseError("formula_without_cached_value",
                                 f"Row {row_no}, column {letters[role]}: the formula has no saved value. "
                                 "Open the file in Excel, recalculate, save and upload again.", row_no)
        code_raw, code_removed = _text(_cell(values, cols["code"]), limits, row_no, letters["code"])
        code = ascii_digits(code_raw).strip()
        name, name_removed = _text(_cell(values, cols["name"]), limits, row_no, letters["name"])
        amounts: dict[str, Decimal | None] = {}
        had_formula = False
        for role in _AMOUNT_ROLES:
            idx = cols.get(role)
            if idx is None:
                continue
            value = _cell(values, idx)
            is_formula = idx < len(formulas) and formulas[idx]
            if is_formula:
                had_formula = True
                if value is None:
                    raise ParseError("formula_without_cached_value",
                                     f"Row {row_no}, column {letters[role]}: the formula has no saved value. "
                                     "Open the file in Excel, recalculate, save and upload again.", row_no)
            amounts[role] = parse_cell(value, row_no, letters[role], decimal_separator)

        if not code:
            label = normalize_for_matching(name)
            if label in TOTAL_ROW or label.startswith(("total ", "grand total", "المجموع", "الاجمالي", "اجمالي")):
                footer = amounts
                continue
            if all(a is None for a in amounts.values()):
                warnings.add("section_headings_skipped", "Rows with a label but no account code or amounts were skipped.",
                             row_no)
                continue
            raise ParseError("missing_account_code", f"Row {row_no} has amounts but no account code.", row_no)

        if len(code) > _MAX_CODE:
            raise ParseError("code_too_long", f"Row {row_no}: the account code is longer than {_MAX_CODE} characters.",
                             row_no)
        if not name:
            name = code
            warnings.add("name_missing", "Accounts without a name were given their code as the name.", row_no)
        if len(name) > _MAX_NAME:
            raise ParseError("name_too_long", f"Row {row_no}: the account name is longer than {_MAX_NAME} characters.",
                             row_no)
        if code_removed or name_removed:
            warnings.add("invisible_characters_removed",
                         "Invisible direction-control or zero-width characters were removed from codes or names.",
                         row_no)
        if looks_like_formula(name) or looks_like_formula(code):
            warnings.add("formula_like_text",
                         "Some names start with =, +, - or @. They are stored as text and escaped on export.", row_no)
        if code in seen_codes:
            warnings.add("duplicate_account_codes", "The same account code appears on more than one row.", row_no)
        seen_codes.setdefault(code, row_no)

        for role in ("period_debit", "period_credit", "closing_debit", "closing_credit", "opening_debit",
                     "opening_credit"):
            v = amounts.get(role)
            if v is not None and v < 0:
                raise ParseError("negative_in_side_column",
                                 f"Row {row_no}, column {letters[role]}: debit and credit columns must not be negative.",
                                 row_no)

        opening = amounts.get("opening")
        if "opening_debit" in cols:
            od, oc = amounts.get("opening_debit"), amounts.get("opening_credit")
            opening = None if od is None and oc is None else (od or Decimal(0)) - (oc or Decimal(0))
        pd, pc = amounts.get("period_debit"), amounts.get("period_credit")
        if layout.kind == "closing_signed":
            closing = amounts.get("closing") or Decimal(0)
        elif layout.kind in ("closing_dr_cr", "dr_cr_as_closing"):
            closing = (amounts.get("closing_debit") or Decimal(0)) - (amounts.get("closing_credit") or Decimal(0))
        else:  # opening_plus_movements
            closing = (opening or Decimal(0)) + (pd or Decimal(0)) - (pc or Decimal(0))
        if had_formula:
            warnings.add("formula_cells", "Some amounts were formulas; their saved values were used (never recalculated).",
                         row_no)

        lines.append(Line(len(lines) + 1, row_no, code, name, normalize_for_matching(name) or normalize_for_matching(code),
                          opening, pd, pc, closing + 0, had_formula))

    if not lines:
        raise ParseError("empty", "No account lines were found below the header row.")

    sum_debit = sum((l.closing for l in lines if l.closing > 0), Decimal(0))
    sum_credit = -sum((l.closing for l in lines if l.closing < 0), Decimal(0))
    net = sum_debit - sum_credit
    if net != 0:
        warnings.add("unbalanced", "Debit and credit balances differ. The TB can be mapped but not locked until "
                                   "the difference is booked to an explicit suspense line.")
    if footer is not None:
        _check_footer(footer, layout, sum_debit, sum_credit, net, warnings)

    return {
        "parser_version": PARSER_VERSION,
        "format": fmt_name,
        "sheet": sheet_name,
        "header_row": layout.header_row,
        "layout": layout.kind,
        "columns": letters,
        "lines": [l.to_dict() for l in lines],
        "control": {"line_count": len(lines), "sum_debit": fmt(sum_debit), "sum_credit": fmt(sum_credit),
                    "net": fmt(net)},
        "warnings": warnings.to_list(),
    }


def _check_footer(footer: dict[str, Decimal | None], layout: Layout, dr: Decimal, cr: Decimal, net: Decimal,
                  warnings: Warnings) -> None:
    fd, fc = footer.get("closing_debit"), footer.get("closing_credit")
    mismatch = False
    if fd is not None or fc is not None:
        mismatch = (fd is not None and fd != dr) or (fc is not None and fc != cr)
    elif footer.get("closing") is not None:
        mismatch = footer["closing"] != net
    if mismatch:
        warnings.add("footer_total_mismatch", "The file's own total row does not agree with the sum of its lines.")


def column_letter(idx: int) -> str:  # re-exported for tests
    return get_column_letter(idx + 1)
